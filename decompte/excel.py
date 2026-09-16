"""Remplissage du modèle Excel « Décompte DGEO » (course d'école ou camp).

On part du modèle fourni (.xltx), on conserve ses formules (G11, règle de trois,
MROUND du total) et on ajuste le nombre de lignes : seules les rubriques utilisées
apparaissent, comme dans les décomptes établis à la main.
"""
from __future__ import annotations

import datetime as dt
import io
import re
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.worksheet import Worksheet

from .models import Dossier

TEMPLATES_DIR = Path(__file__).parent / "templates"
TEMPLATES = {"course": "modele_course.xltx", "camp": "modele_camp.xltx"}

FIRST_DATA_ROW = 12
STYLE_COLS = range(1, 12)  # A..K
DATE_FMT = "dd.mm.yyyy"


def _find_total_row(ws: Worksheet) -> int:
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if isinstance(v, str) and "total du remboursement" in v.lower():
            return r
    raise ValueError("Ligne « Total du remboursement » introuvable dans le modèle")


def _find_date_label_row(ws: Worksheet, after: int) -> int | None:
    for r in range(after, ws.max_row + 1):
        for c in range(1, 12):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().lower().startswith("date"):
                return r
    return None


def _copy_row_style(ws: Worksheet, src_row: int, dst_row: int) -> None:
    for c in STYLE_COLS:
        s, d = ws.cell(src_row, c), ws.cell(dst_row, c)
        d.font = copy(s.font)
        d.border = copy(s.border)
        d.fill = copy(s.fill)
        d.alignment = copy(s.alignment)
        d.number_format = s.number_format
        d.protection = copy(s.protection)


def _resize_data_rows(ws: Worksheet, n_rows: int) -> int:
    """Ajuste le nombre de lignes de données (12..total-1) à n_rows. Retourne la nouvelle ligne total."""
    total_row = _find_total_row(ws)
    slots = total_row - FIRST_DATA_ROW
    n_rows = max(1, n_rows)
    delta = n_rows - slots
    if delta == 0:
        return total_row
    # fusions et hauteurs situées sous la zone de données : à décaler à la main
    merged_below = [r for r in list(ws.merged_cells.ranges) if r.min_row >= total_row]
    for r in merged_below:
        ws.unmerge_cells(str(r))
    heights = {r: ws.row_dimensions[r].height for r in range(total_row, ws.max_row + 1) if ws.row_dimensions[r].height}
    if delta > 0:
        ws.insert_rows(total_row, delta)
        for r in range(total_row, total_row + delta):
            _copy_row_style(ws, FIRST_DATA_ROW, r)
    else:
        ws.delete_rows(FIRST_DATA_ROW + n_rows, -delta)
    for r in merged_below:
        ws.merge_cells(
            start_row=r.min_row + delta, start_column=r.min_col, end_row=r.max_row + delta, end_column=r.max_col
        )
    for r in range(total_row, total_row + max(delta, 0) + len(heights) + 5):
        ws.row_dimensions[r].height = None
    for r, hgt in heights.items():
        ws.row_dimensions[r + delta].height = hgt
    if ws.print_area:
        area = ws.print_area if isinstance(ws.print_area, str) else ws.print_area[0]
        cells = area.split("!")[-1].replace("$", "")
        start, end = cells.split(":")
        col_end = "".join(ch for ch in end if ch.isalpha())
        row_end = int("".join(ch for ch in end if ch.isdigit())) + delta
        ws.print_area = f"{start}:{col_end}{row_end}"
    return total_row + delta


def _parse_date(s: str | None) -> dt.date | None:
    if not s:
        return None
    try:
        d, m, y = s.split(".")
        return dt.date(int(y), int(m), int(d))
    except (ValueError, AttributeError):
        return None


def _direct_formula(row) -> "str | float":
    """Colonne I d'une ligne « saisie directe » : « =6*2.8+6*4.2+2*2.1 » (le détail des tarifs adultes retenus),
    ou le montant si le détail n'est pas disponible (ligne saisie à la main)."""
    amount = round(row.cout_direct or 0.0, 2)
    detail = (row.formule or "").replace(" ", "")
    if not detail or not re.fullmatch(r"[0-9.*+EURCHF]+", detail):
        return amount
    expr = detail.replace("EUR", "").replace("CHF", "")
    try:
        if abs(round(eval(expr, {"__builtins__": {}}, {}), 2) - amount) > 0.011:  # noqa: S307 (expression numérique contrôlée)
            return amount
    except Exception:  # noqa: BLE001
        return amount
    return f"=ROUND({expr},2)"


def build_workbook(dossier: Dossier, today: dt.date | None = None) -> bytes:
    template = TEMPLATES_DIR / TEMPLATES.get(dossier.type_activite, TEMPLATES["course"])
    wb = openpyxl.load_workbook(template)
    wb.template = False
    ws = wb.worksheets[0]

    # En-tête
    ws["C5"] = dossier.classe or None
    ws["C6"] = dossier.enseignant or None
    if dossier.type_activite == "camp":
        ws["C7"] = f"Camp – {dossier.activite}" if dossier.activite else "Camp"
    else:
        ws["C7"] = "Course d'école"
    d0, d1 = _parse_date(dossier.date_debut), _parse_date(dossier.date_fin)
    if d0 and d1 and d1 != d0:
        ws["C8"] = f"{d0:%d.%m.%Y} – {d1:%d.%m.%Y}"
    elif d0:
        ws["C8"] = d0
        ws["C8"].number_format = DATE_FMT
    else:
        ws["C8"] = None
    ws["J8"] = dossier.numero or None

    # Effectifs (G11 = SUM(D11:F11) reste une formule du modèle)
    eff = dossier.effectifs
    ws["D11"] = eff.non_titres or None
    ws["E11"] = eff.eleves
    ws["F11"] = eff.titres

    # Lignes
    rows = dossier.rows
    total_row = _resize_data_rows(ws, len(rows))
    for i in range(max(1, len(rows))):
        r = FIRST_DATA_ROW + i
        for c in range(2, 11):
            ws.cell(r, c).value = None
    for i, row in enumerate(rows):
        r = FIRST_DATA_ROW + i
        ws.cell(r, 2).value = row.rubrique
        ws.cell(r, 3).value = row.libelle
        if row.mode == "direct":
            # coût total des billets en H (information) ; part État en I sous forme de formule montrant le calcul
            ws.cell(r, 8).value = round(row.cout_total, 2) if row.cout_total is not None else None
            ws.cell(r, 9).value = _direct_formula(row)
        else:
            ws.cell(r, 8).value = round(row.cout_total or 0.0, 2)
            ws.cell(r, 9).value = f"=IFERROR(H{r}/$G$11*$F$11,0)"
        ws.cell(r, 10).value = f"=I{r}"
        for c in (8, 9, 10):
            ws.cell(r, c).number_format = "#,##0.00"
    if not rows:
        r = FIRST_DATA_ROW
        ws.cell(r, 9).value = f"=IFERROR(H{r}/$G$11*$F$11,0)"
        ws.cell(r, 10).value = f"=I{r}"
    last = FIRST_DATA_ROW + max(1, len(rows)) - 1
    ws.cell(total_row, 10).value = f"=MROUND(SUM(J{FIRST_DATA_ROW}:J{last}),0.05)"

    # Date de signature
    date_row = _find_date_label_row(ws, total_row + 1)
    if date_row:
        cell = ws.cell(date_row, 9)
        cell.value = _parse_date(dossier.date_decompte) or today or dt.date.today()
        cell.number_format = DATE_FMT

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def output_filename(dossier: Dossier) -> str:
    base = dossier.numero or Path(dossier.filename).stem or "decompte"
    base = "".join(ch for ch in base if ch.isalnum() or ch in "-_")
    return f"{base or 'decompte'}.xlsx"
